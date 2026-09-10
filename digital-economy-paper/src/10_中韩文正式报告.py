from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "work" / "formal_final_20260816"
FINAL = ROOT / "outputs" / "韩国数字化战略论文_正式版_20260816"
FIGURES = FINAL / "02_模型与结果" / "图表"
INPUTS = json.loads((WORK / "formal_inputs.json").read_text(encoding="utf-8"))

PAGE_WIDTH_DXA = 12240
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
NAVY = "1F4D78"
TEXT = "1F2933"
MUTED = "5C6773"
LIGHT = "F4F6F9"
LINE = "D6DEE6"
WARNING = "FFF4DE"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_borders(table, color=LINE, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_run_font(run, font: str, east_asia: str, size=None, bold=None, italic=None, color=TEXT) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph_border_bottom(paragraph, color=BLUE, size="10", space="4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph, prefix: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(prefix)
    set_run_font(run, "Calibri", "Microsoft YaHei", 9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_bullet(doc: Document, text: str, font: str, east: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, font, east, 11)


def add_number(doc: Document, text: str, font: str, east: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, font, east, 11)


def add_para(doc: Document, text: str, font: str, east: str, bold=False, italic=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8, before=0) -> Any:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    run = p.add_run(text)
    set_run_font(run, font, east, 11, bold=bold, italic=italic, color=color)
    return p


def add_callout(doc: Document, label: str, text: str, font: str, east: str, fill=LIGHT) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="C6D4E0", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r1 = p.add_run(f"{label} ")
    set_run_font(r1, font, east, 11, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, font, east, 11, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[int], font: str, east: str, font_size=9.2) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(header))
        set_run_font(run, font, east, font_size, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            if row_index % 2 == 1:
                set_cell_shading(cells[col_index], LIGHT)
            p = cells[col_index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run("" if value is None else str(value))
            set_run_font(run, font, east, font_size, color=TEXT)
    set_table_geometry(table, widths)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def fmt_coef(row: dict) -> str:
    return f"{row['coefficient']:.6f}"


def fmt_p(value: float | None) -> str:
    if value is None:
        return ""
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def outcome_name(outcome: str, lang: str) -> str:
    cn = {"SalesGrowth_w_t1": "次年销售增长率", "ROA_AvgAssets_w_t1": "次年ROA", "OperatingMargin_w_t1": "次年营业利润率", "AlignedCAR_0_p2_w": "CAR[0,+2]"}
    kr = {"SalesGrowth_w_t1": "차기 매출성장률", "ROA_AvgAssets_w_t1": "차기 ROA", "OperatingMargin_w_t1": "차기 영업이익률", "AlignedCAR_0_p2_w": "CAR[0,+2]"}
    return (cn if lang == "cn" else kr).get(outcome, outcome)


def setup_document(lang: str) -> tuple[Document, str, str]:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True
    font = "Calibri"
    east = "Microsoft YaHei" if lang == "cn" else "Malgun Gothic"

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ]:
        style = styles[name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        style.font.size = Pt(11)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    htext = "韩国上市公司数字化战略研究 | 正式实证报告" if lang == "cn" else "한국 상장기업 디지털 전략 연구 | 공식 실증보고서"
    hr = hp.add_run(htext)
    set_run_font(hr, font, east, 9, bold=True, color=MUTED)
    paragraph_border_bottom(hp, color=LINE, size="4", space="3")
    footer = section.footer
    add_page_number(footer.paragraphs[0], "第 " if lang == "cn" else "p. ")
    return doc, font, east


def add_cover(doc: Document, lang: str, font: str, east: str) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = "正式实证分析报告" if lang == "cn" else "공식 실증분석 보고서"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(kicker)
    set_run_font(r, font, east, 11, bold=True, color=BLUE)
    title_cn = "企业数字化战略的“双轮驱动”效应"
    sub_cn = "内部商业模式重构与外部资本市场信号：来自韩国上市公司的证据"
    title_kr = "기업 디지털 전략의 이중 구동 효과"
    sub_kr = "내부 비즈니스 모델 재구성과 외부 자본시장 신호: 한국 상장기업의 증거"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title_cn if lang == "cn" else title_kr)
    set_run_font(r, font, east, 26, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run(sub_cn if lang == "cn" else sub_kr)
    set_run_font(r, font, east, 14, color=BLUE)
    p = doc.add_paragraph()
    paragraph_border_bottom(p, color=BLUE, size="10", space="6")
    for _ in range(3):
        doc.add_paragraph()
    meta = [
        ("样本", f"{INPUTS['sample']['firms']:,}家公司，{INPUTS['sample']['rows']:,}个公司年度"),
        ("期间", f"{INPUTS['sample']['years'][0]}-{INPUTS['sample']['years'][1]}财年"),
        ("数据", "DART韩文年报、OpenDART财务数据、韩国股票与市场指数"),
        ("版本", "冻结数据与模型审计正式版，2026-08-16"),
    ] if lang == "cn" else [
        ("표본", f"{INPUTS['sample']['firms']:,}개 기업, {INPUTS['sample']['rows']:,}개 기업-연도"),
        ("기간", f"{INPUTS['sample']['years'][0]}-{INPUTS['sample']['years'][1]} 회계연도"),
        ("자료", "DART 한국어 사업보고서, OpenDART 재무자료, 한국 주가 및 시장지수"),
        ("버전", "동결자료 및 모형감사 공식본, 2026-08-16"),
    ]
    add_table(doc, [meta[0][0], meta[0][1]], [[x[0], x[1]] for x in meta[1:]], [2000, 7360], font, east, 10)
    doc.add_page_break()


def add_contents(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("报告结构" if lang == "cn" else "보고서 구성", level=1)
    sections_cn = ["执行摘要", "研究问题与理论主线", "数据与样本构建", "句子级数字化战略测量", "变量与模型", "内部经营结果", "资本市场结果", "并发公告审计", "非线性、机制与稳健性", "原因解释：为何严格模型不显著", "韩国制度环境与外生冲击", "结论、贡献与限制", "复现与文件说明", "参考来源"]
    sections_kr = ["요약", "연구질문과 이론적 논리", "자료와 표본구축", "문장수준 디지털 전략 측정", "변수와 모형", "내부 경영성과", "자본시장 반응", "동시공시 감사", "비선형성·메커니즘·강건성", "원인 해석: 엄격한 모형에서 유의성이 사라지는 이유", "한국 제도환경과 외생충격", "결론·기여·한계", "재현 및 파일안내", "참고자료"]
    for item in sections_cn if lang == "cn" else sections_kr:
        add_bullet(doc, item, font, east)
    add_callout(
        doc,
        "阅读提示" if lang == "cn" else "해석 안내",
        "本报告把确认性模型、探索性模型与尚待数据支持的因果扩展明确分开。统计不显著不等于不存在任何经济机制，但不能被改写为支持性证据。" if lang == "cn" else "본 보고서는 확인적 모형, 탐색적 모형, 추가자료가 필요한 인과확장을 구분한다. 통계적 비유의성은 모든 경제적 메커니즘의 부재를 뜻하지 않지만 지지증거로 재해석할 수 없다.",
        font,
        east,
    )


def add_executive_summary(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("1. 执行摘要" if lang == "cn" else "1. 요약", level=1)
    preferred = next(x for x in INPUTS["car_core"] if x["model_id"] == "CAR_EVENT_DATE_FE_TWO_WAY")
    if lang == "cn":
        add_para(doc, f"本研究以韩国上市公司为对象，将企业数字化战略从简单词频重新定义为年报中可识别的实质性数字行动。冻结面板包括{INPUTS['sample']['firms']:,}家公司、{INPUTS['sample']['rows']:,}个公司年度，覆盖{INPUTS['sample']['years'][0]}-{INPUTS['sample']['years'][1]}财年。核心解释变量为实质性数字行动句占年报去重总句数的比例，并进一步区分商业模式重构、运营数字化、纯信号与口号式披露。", font, east)
        add_para(doc, "内部路径以次年销售增长、次年ROA和次年营业利润率衡量，采用公司和年份固定效应并按公司聚类。三项确认性结果均不显著，且替代变量、零值处理、公司趋势、一阶差分、扩展控制和文本时间边界限制没有形成一致的正向或非线性证据。", font, east)
        add_para(doc, f"外部路径以严格对齐披露交易日的CAR[0,+2]衡量。首选的公司固定效应、事件日固定效应及公司/事件日双向聚类模型得到系数{preferred['coefficient']:.6f}（标准误{preferred['std_error']:.6f}，p={preferred['p_value']:.3f}，FDR={preferred['fdr_within_family']:.3f}）。点估计对应数字化战略提高1个标准差时三日CAR低约0.269个百分点，但它只是一种边缘性的负向关联。滞后控制结果不显著，未来数字化水平安慰剂显著，故不能解释为已识别的资本市场惩罚。", font, east)
        add_callout(doc, "正式结论", "韩国上市公司年报中的实质数字行动披露没有显示可稳健归因的短期经营改善；资本市场对更强披露呈现弱负向定价倾向，但多重检验、模型敏感性和安慰剂结果要求保留判断。论文最有价值的贡献是更严格的文本测量、时间对齐和对‘披露不等于实现’的边界识别。", font, east, fill=WARNING)
    else:
        add_para(doc, f"본 연구는 한국 상장기업을 대상으로 기업의 디지털 전략을 단순 키워드 빈도가 아니라 사업보고서에서 식별되는 실질적 디지털 행동으로 재정의한다. 동결 패널은 {INPUTS['sample']['firms']:,}개 기업, {INPUTS['sample']['rows']:,}개 기업-연도, {INPUTS['sample']['years'][0]}-{INPUTS['sample']['years'][1]} 회계연도를 포함한다. 핵심 설명변수는 중복제거 총문장 중 실질 디지털 행동문장의 비율이며 비즈니스모델 재구성, 운영 디지털화, 순수신호와 상투적 공시를 추가로 구분한다.", font, east)
        add_para(doc, "내부 경로는 차기 매출성장률, 차기 ROA와 차기 영업이익률로 측정하고 기업 및 연도 고정효과와 기업 군집표준오차를 사용한다. 세 확인적 결과는 모두 유의하지 않았으며 대체 측정, 영(0)값 처리, 기업추세, 1차 차분, 확장 통제와 텍스트 시점 제한에서도 일관된 양(+) 또는 비선형 증거가 나타나지 않았다.", font, east)
        add_para(doc, f"외부 경로는 공시거래일에 엄격히 정렬한 CAR[0,+2]로 측정한다. 기업 고정효과, 사건일 고정효과 및 기업/사건일 이중 군집을 적용한 선호모형의 계수는 {preferred['coefficient']:.6f}(표준오차 {preferred['std_error']:.6f}, p={preferred['p_value']:.3f}, FDR={preferred['fdr_within_family']:.3f})이다. 디지털 전략 1표준편차 증가는 약 0.269%p 낮은 3일 CAR과 연관되지만 이는 경계적인 음(-)의 상관이다. 시차 통제모형은 유의하지 않고 미래수준 위약검정이 실패하므로 식별된 시장페널티로 해석할 수 없다.", font, east)
        add_callout(doc, "공식 결론", "실질 디지털 행동 공시는 단기 경영성과 개선과 강건하게 연결되지 않았다. 자본시장은 강한 공시에 약한 음(-)의 가격반응을 보일 가능성이 있으나 다중검정, 모형민감성과 위약검정 때문에 신중한 해석이 필요하다. 엄격한 텍스트 측정과 시점정렬, 그리고 ‘공시와 실행의 구분’이 본 연구의 핵심 기여이다.", font, east, fill=WARNING)


def add_theory(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("2. 研究问题与理论主线" if lang == "cn" else "2. 연구질문과 이론적 논리", level=1)
    if lang == "cn":
        add_para(doc, "研究主线是同一项数字化战略如何通过两个不同的时间和评价机制发生作用。内部轮指组织能力、流程和商业模式重构，理论上需要互补资产、学习和实施周期才能转化为经营绩效。外部轮指资本市场从年报中观察数字行动并及时调整预期，但投资者同时评估实施成本、兑现概率、信息复杂性和管理层宣传动机。", font, east)
        add_para(doc, "由此，内部效应不必与公告日市场反应同号或同时出现。短期经营结果可能因调整成本、项目周期和收益确认滞后而接近零；公告日反应则可能为负，因为更强的数字化披露既传递成长机会，也暴露资本支出、组织变革与执行不确定性。该逻辑允许‘内部短期不显著+外部弱负向’成为理论上连贯而非失败的结果，但仍须尊重统计证据边界。", font, east)
        add_bullet(doc, "研究问题1：实质性数字化战略披露是否预测下一期经营绩效？", font, east)
        add_bullet(doc, "研究问题2：商业模式重构是否改变或构成内部作用路径？", font, east)
        add_bullet(doc, "研究问题3：资本市场是否在年报披露窗口对数字化战略作出定价反应？", font, east)
        add_bullet(doc, "研究问题4：关系是否具有非线性、市场异质性或受口号式信号干扰？", font, east)
    else:
        add_para(doc, "연구의 중심 논리는 하나의 디지털 전략이 서로 다른 시간과 평가기제를 통해 두 경로로 작동한다는 것이다. 내부 경로는 조직역량, 프로세스와 비즈니스모델 재구성으로서 보완자산, 학습과 실행기간을 거쳐 경영성과로 전환된다. 외부 경로에서는 투자자가 사업보고서의 디지털 행동을 즉시 관찰하지만 동시에 실행비용, 실현확률, 정보복잡성과 경영자 홍보동기를 평가한다.", font, east)
        add_para(doc, "따라서 내부효과와 공시일 시장반응은 동일한 부호나 시점에 나타날 필요가 없다. 단기 경영성과는 전환비용과 사업주기 때문에 0에 가까울 수 있고 공시일 반응은 성장기회와 함께 자본지출·조직변화·실행불확실성을 반영하여 음(-)일 수 있다. 이 논리는 ‘내부 단기 비유의+외부 약한 음(-)’을 이론적으로 연결하지만 통계적 증거의 한계를 넘어서지는 않는다.", font, east)
        add_bullet(doc, "연구질문 1: 실질적 디지털 전략 공시는 차기 경영성과를 예측하는가?", font, east)
        add_bullet(doc, "연구질문 2: 비즈니스모델 재구성은 내부 작동경로를 구성하거나 조절하는가?", font, east)
        add_bullet(doc, "연구질문 3: 자본시장은 연차보고서 공시창에서 디지털 전략을 가격에 반영하는가?", font, east)
        add_bullet(doc, "연구질문 4: 비선형성, 시장 이질성 또는 상투적 신호가 관계를 변화시키는가?", font, east)


def add_data_measurement(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("3. 数据与样本构建" if lang == "cn" else "3. 자료와 표본구축", level=1)
    if lang == "cn":
        for text in [
            "公司母表来自DART/OpenDART，使用稳定corp_code连接公司名称、股票代码、市场分类和财年。年报文本以韩文DART为主，因为英文年报存在严重选择性与长期缺失。财务报表优先使用合并口径，并按公司和财年连接到文本记录。",
            "股票结果以报告文本实际可观察的披露日期为准，而不是仅用财年末或原始年报日期。对修订文本，事件日与所使用文本版本一致；非交易日顺延到下一交易日。市场模型估计窗和事件窗均从缓存的个股与KOSPI/KOSDAQ指数日行情构造。",
            f"冻结主面板有{INPUTS['sample']['rows']:,}个公司年度、{INPUTS['sample']['firms']:,}家公司，重复公司年度为{INPUTS['sample']['duplicate_firm_years']}。数字化战略为零的公司年度占{100*INPUTS['sample']['digital_zero_share']:.2f}%，说明零膨胀与公司内高持续性是测量和识别的真实约束。",
        ]:
            add_para(doc, text, font, east)
    else:
        for text in [
            "기업 모표는 DART/OpenDART에서 구축하고 안정적인 corp_code로 기업명, 종목코드, 시장구분과 회계연도를 연결하였다. 영문 연차보고서는 선택성과 장기 결측이 커서 본문 텍스트는 한국어 DART 사업보고서를 우선 사용하였다. 재무제표는 연결기준을 우선하며 기업-회계연도로 텍스트와 결합하였다.",
            "주가분석의 사건일은 회계연도 말이나 최초 보고일이 아니라 실제 분석한 텍스트가 시장에서 관찰된 공시일이다. 정정텍스트는 사용한 문서버전과 사건일을 일치시키고 비거래일은 다음 거래일로 이동하였다. 시장모형 추정창과 사건창은 개별주식 및 KOSPI/KOSDAQ 지수 일별자료로 구성하였다.",
            f"동결 주패널은 {INPUTS['sample']['rows']:,}개 기업-연도와 {INPUTS['sample']['firms']:,}개 기업으로 구성되며 중복 기업-연도는 {INPUTS['sample']['duplicate_firm_years']}개이다. 디지털 전략이 0인 기업-연도 비중은 {100*INPUTS['sample']['digital_zero_share']:.2f}%로 영(0)값 집중과 높은 기업내 지속성이 실질적 측정·식별 제약임을 보여준다.",
        ]:
            add_para(doc, text, font, east)

    year_rows = [[x["year"], f"{x['observations']:,}", f"{x['firms']:,}", f"{100*x['digital_mean']:.3f}%", f"{100*x['positive_digital_share']:.1f}%"] for x in INPUTS["sample"]["year_counts"]]
    add_table(doc, ["年份" if lang == "cn" else "연도", "公司年度" if lang == "cn" else "기업-연도", "公司数" if lang == "cn" else "기업수", "平均DigitalStrategy" if lang == "cn" else "평균 DigitalStrategy", "正值占比" if lang == "cn" else "양(+)값 비중"], year_rows, [1000, 1700, 1600, 2600, 2460], font, east, 8.8)
    if (FIGURES / "01_数字化战略年度趋势.png").exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(FIGURES / "01_数字化战略年度趋势.png"), width=Inches(6.3))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = caption.add_run("图1 数字化战略年度趋势" if lang == "cn" else "그림 1 디지털 전략의 연도별 추세")
        set_run_font(rr, font, east, 9, italic=True, color=MUTED)

    doc.add_heading("4. 句子级数字化战略测量" if lang == "cn" else "4. 문장수준 디지털 전략 측정", level=1)
    if lang == "cn":
        add_para(doc, "测量流程先进行年报正文提取、表格与导航噪声处理、句子切分和去重，再识别数字技术实体与行动谓词。只有同时包含数字技术或数字业务对象，并表达投资、建设、上线、整合、优化、推出、转型渠道或收入模式等可观察行动的句子，才进入实质性数字战略分子。仅表达重要性、趋势、愿景或对数字时代的泛化回应，归为口号式披露或低权重候选。", font, east)
        add_para(doc, "理论依据是战略必须包含资源承诺、组织选择或商业活动改变，而非仅出现技术名词。句子级联合条件降低了digital、AI、data、platform随时间普遍化造成的测量误差，也允许把商业模式重构与运营技术部署分开。", font, east)
    else:
        add_para(doc, "측정과정은 사업보고서 본문 추출, 표·탐색메뉴 잡음 제거, 문장분할과 중복제거 후 디지털 기술대상과 행동술어를 함께 식별한다. 디지털 기술 또는 디지털 사업대상을 포함하면서 투자, 구축, 도입, 통합, 최적화, 출시, 채널 또는 수익모델 전환과 같은 관찰 가능한 행동을 표현한 문장만 실질 디지털 전략의 분자에 포함한다. 중요성, 추세, 비전 또는 디지털 시대에 대한 일반적 대응만 표현한 문장은 상투적 공시 또는 저가중 후보로 분류한다.", font, east)
        add_para(doc, "이론적 근거는 전략이 기술명사의 출현이 아니라 자원투입, 조직적 선택 또는 사업활동의 변화를 포함해야 한다는 점이다. 문장수준 결합조건은 digital, AI, data, platform 용어의 전반적 확산에 따른 측정오차를 줄이고 비즈니스모델 재구성과 운영기술 배치를 구분한다.", font, east)

    measure_rows_cn = [
        ["DigitalStrategy", "实质数字行动句 / 年报总句数", "主变量"],
        ["BMI_Reconstruction", "数字商业模式重构句 / 总句数", "内部路径"],
        ["OperationalDigital", "max(实质行动句-BMI句,0) / 总句数", "运营成分"],
        ["PureSignal", "max(市场可见数字句-实质行动句,0) / 总句数", "净信号成分"],
        ["RelativeDigitalSignal", "DigitalStrategy-行业×年份均值", "共同趋势调整"],
        ["Boilerplate", "口号式数字句 / 总句数", "效度诊断"],
    ]
    measure_rows_kr = [
        ["DigitalStrategy", "실질 디지털 행동문장 / 총문장", "주변수"],
        ["BMI_Reconstruction", "디지털 비즈니스모델 재구성문장 / 총문장", "내부 경로"],
        ["OperationalDigital", "max(실질행동-BMI문장,0) / 총문장", "운영 성분"],
        ["PureSignal", "max(시장가시 디지털문장-실질행동,0) / 총문장", "순신호 성분"],
        ["RelativeDigitalSignal", "DigitalStrategy-산업×연도 평균", "공통추세 조정"],
        ["Boilerplate", "상투적 디지털문장 / 총문장", "타당성 진단"],
    ]
    add_table(doc, ["变量" if lang == "cn" else "변수", "定义" if lang == "cn" else "정의", "用途" if lang == "cn" else "용도"], measure_rows_cn if lang == "cn" else measure_rows_kr, [2300, 4660, 2400], font, east, 9)
    add_callout(doc, "效度边界" if lang == "cn" else "타당성 한계", "MarketSignalDisclosure与DigitalStrategy在现有规则下相关性接近1，且94.25%的公司年度两类句数相同，因此正式模型不把它当作独立外部机制。改用PureSignal或SignalGap可以减少机械重叠。真正的第二编码者留出验证仍应在投稿前用独立判断完成；本正式包不把未经独立完成的一致性数字当作证据。" if lang == "cn" else "현 규칙에서 MarketSignalDisclosure와 DigitalStrategy의 상관이 거의 1이고 94.25%의 기업-연도에서 두 문장수가 동일하다. 따라서 이를 독립적인 외부 메커니즘으로 사용하지 않고 PureSignal 또는 SignalGap으로 기계적 중복을 줄인다. 독립 제2코더 홀드아웃 검증은 투고 전 실제 독립판단으로 완료해야 하며, 본 공식본은 독립적으로 완료되지 않은 일치도 수치를 증거로 사용하지 않는다.", font, east, fill=WARNING)


def add_models(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("5. 变量与模型" if lang == "cn" else "5. 변수와 모형", level=1)
    doc.add_heading("5.1 内部经营模型" if lang == "cn" else "5.1 내부 경영모형", level=2)
    equation = "Y(i,t+1) = β·DigitalStrategy(i,t) + γ'X(i,t) + α(i) + λ(t) + ε(i,t)"
    add_callout(doc, "模型" if lang == "cn" else "모형", equation, font, east)
    if lang == "cn":
        add_para(doc, "因变量分别为次年销售增长率、次年ROA和次年营业利润率。核心控制变量包括公司规模、杠杆、同期ROA、销售增长、流动性、营业利润率、公司年龄、亏损虚拟变量、负净资产虚拟变量和年报篇幅。扩展模型加入资产周转率、营运资本比率、员工规模和最大股东及关联方持股；2020年以后另加入研发强度与外部董事比例。连续变量按1%和99%缩尾。", font, east)
        add_para(doc, "公司固定效应吸收不随时间变化的管理质量、行业定位和企业文化，年份固定效应吸收宏观经济与共同数字化趋势，标准误按公司聚类以允许公司内序列相关。模型识别来自同一公司跨年度数字战略强度的变化，因此指标高持久性会降低有效信息和检验力。", font, east)
    else:
        add_para(doc, "종속변수는 차기 매출성장률, 차기 ROA와 차기 영업이익률이다. 핵심 통제변수는 기업규모, 레버리지, 당기 ROA, 매출성장, 유동성, 영업이익률, 기업연령, 손실더미, 자본잠식더미와 보고서 길이이다. 확장모형은 자산회전율, 운전자본비율, 종업원규모와 최대주주 및 특수관계인 지분을 추가하며 2020년 이후에는 연구개발집약도와 사외이사비율도 검토한다. 연속변수는 1%와 99%에서 윈저라이징하였다.", font, east)
        add_para(doc, "기업 고정효과는 시간불변 경영품질, 산업포지셔닝과 기업문화를 흡수하고 연도 고정효과는 거시경제와 공통 디지털 추세를 통제한다. 기업 군집표준오차는 기업내 시계열 상관을 허용한다. 식별은 동일기업의 연도별 디지털 전략 변화에서 나오므로 높은 지속성은 유효정보와 검정력을 낮춘다.", font, east)

    doc.add_heading("5.2 资本市场模型" if lang == "cn" else "5.2 자본시장 모형", level=2)
    equation = "CAR(i,e)[0,+2] = β·DigitalStrategy(i,t) + γ'X(i,t) + α(i) + δ(e) + u(i,e)"
    add_callout(doc, "模型" if lang == "cn" else "모형", equation, font, east)
    if lang == "cn":
        add_para(doc, "首选模型使用公司固定效应和精确披露交易日固定效应，标准误按公司与披露日双向聚类。事件日固定效应比一般年份固定效应更合适，因为大量韩国年报集中在法定披露期限附近，同日市场冲击可能造成跨公司相关。控制变量为核心财务变量、年报篇幅、财年股票收益和日收益波动率；滞后控制模型用于降低同时控制披露结果的风险。", font, east)
    else:
        add_para(doc, "선호모형은 기업 고정효과와 정확한 공시거래일 고정효과를 사용하고 기업 및 공시일에 이중 군집표준오차를 적용한다. 한국 사업보고서는 법정기한 부근에 집중되므로 동일일 시장충격이 기업간 상관을 만들 수 있어 사건일 고정효과가 일반 연도효과보다 적합하다. 통제변수는 핵심 재무변수, 보고서 길이, 회계연도 주식수익률과 일별 변동성이며 시차 통제모형으로 동시적 공시결과 통제위험을 점검한다.", font, east)

    doc.add_heading("5.3 非线性与交互项" if lang == "cn" else "5.3 비선형성과 상호작용", level=2)
    if lang == "cn":
        add_para(doc, "平方项模型并不天然更准确。只有在理论上存在阈值、调整成本或倒U形机制，并且数据在公司内提供足够曲率信息时，非线性才有优势。本研究把平方项、正值样本三次项、四分位组和两部分模型作为探索性检验，并用联合显著性而非单个项判断曲线。Digital×BMI和Digital×PureSignal用于检验条件性差异，不被表述为因果中介。", font, east)
    else:
        add_para(doc, "제곱항 모형이 자동으로 더 정확한 것은 아니다. 이론적으로 임계점, 전환비용 또는 역U자 메커니즘이 존재하고 기업내 자료가 충분한 곡률정보를 제공할 때만 비선형성이 유리하다. 본 연구는 제곱항, 양(+)값 표본의 3차항, 분위집단과 허들모형을 탐색적으로 사용하고 단일항이 아니라 결합유의성으로 곡선을 판단한다. Digital×BMI와 Digital×PureSignal은 조건부 차이를 검토하지만 인과적 매개효과로 표현하지 않는다.", font, east)


def add_results(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("6. 内部经营结果" if lang == "cn" else "6. 내부 경영성과", level=1)
    rows = [[outcome_name(x["outcome"], lang), fmt_coef(x), f"{x['std_error']:.6f}", fmt_p(x["p_value"]), fmt_p(x["fdr_within_family"]), f"{x['n']:,}"] for x in INPUTS["operating_core"]]
    add_table(doc, ["因变量" if lang == "cn" else "종속변수", "系数" if lang == "cn" else "계수", "标准误" if lang == "cn" else "표준오차", "p值" if lang == "cn" else "p값", "FDR", "N"], rows, [2600, 1350, 1350, 1100, 1100, 1860], font, east, 9)
    if lang == "cn":
        add_para(doc, "销售增长的系数接近零，ROA与营业利润率的点估计为负，但均不显著。完整经营模型族包含83个数字化核心系数，没有原始p<0.05，也没有FDR<0.05。文本修订限制为未修订年报、修订滞后不超过90天或275天时，12项核心结果同样全部不显著。", font, east)
        add_para(doc, "这组结果不应解释为数字化没有价值。更准确的解释是：年报中实质行动的相对披露强度，无法在公司内年度变化上稳健预测下一年的短期会计绩效。可能原因包括实施收益滞后、转型成本先行、不同项目质量高度异质、年报语言对实际投入强度的测量仍不充分，以及公司固定效应下可利用变化偏小。", font, east)
    else:
        add_para(doc, "매출성장 계수는 0에 가깝고 ROA와 영업이익률의 점추정치는 음(-)이지만 모두 유의하지 않다. 전체 경영모형군의 디지털 핵심계수 83개 중 원 p<0.05 및 FDR<0.05는 모두 0개이다. 미정정 보고서, 정정시차 90일 이하 또는 275일 이하로 제한한 12개 핵심결과도 모두 비유의적이다.", font, east)
        add_para(doc, "이는 디지털화가 가치가 없음을 뜻하지 않는다. 보다 정확한 해석은 사업보고서에 나타난 실질행동의 상대적 공시강도가 기업내 연도변화에서 다음해 단기 회계성과를 강건하게 예측하지 못한다는 것이다. 실행효과의 시차, 선행 전환비용, 사업품질의 이질성, 언어측정과 실제투자강도의 차이, 고정효과에서 활용 가능한 변화의 부족이 가능한 이유이다.", font, east)

    doc.add_heading("7. 资本市场结果" if lang == "cn" else "7. 자본시장 반응", level=1)
    ids = ["CAR_YEAR_FE_FIRM_CLUSTER", "CAR_EVENT_DATE_FE_FIRM_CLUSTER", "CAR_EVENT_DATE_FE_TWO_WAY", "CAR_EVENT_DATE_FE_LAGGED_CONTROLS", "CAR_MARKET_DATE_FE_TWO_WAY", "CAR_SAMPLE_KOSPI", "CAR_SAMPLE_KOSDAQ"]
    labels_cn = {"CAR_YEAR_FE_FIRM_CLUSTER": "年份FE/公司聚类", "CAR_EVENT_DATE_FE_FIRM_CLUSTER": "事件日FE/公司聚类", "CAR_EVENT_DATE_FE_TWO_WAY": "事件日FE/双向聚类（首选）", "CAR_EVENT_DATE_FE_LAGGED_CONTROLS": "滞后控制", "CAR_MARKET_DATE_FE_TWO_WAY": "市场×事件日FE", "CAR_SAMPLE_KOSPI": "KOSPI", "CAR_SAMPLE_KOSDAQ": "KOSDAQ"}
    labels_kr = {"CAR_YEAR_FE_FIRM_CLUSTER": "연도FE/기업군집", "CAR_EVENT_DATE_FE_FIRM_CLUSTER": "사건일FE/기업군집", "CAR_EVENT_DATE_FE_TWO_WAY": "사건일FE/이중군집(선호)", "CAR_EVENT_DATE_FE_LAGGED_CONTROLS": "시차 통제", "CAR_MARKET_DATE_FE_TWO_WAY": "시장×사건일FE", "CAR_SAMPLE_KOSPI": "KOSPI", "CAR_SAMPLE_KOSDAQ": "KOSDAQ"}
    lookup = {x["model_id"]: x for x in INPUTS["car_core"]}
    car_rows = []
    for model_id in ids:
        x = lookup[model_id]
        car_rows.append([(labels_cn if lang == "cn" else labels_kr)[model_id], fmt_coef(x), f"{x['std_error']:.6f}", fmt_p(x["p_value"]), fmt_p(x["fdr_within_family"]), f"{x['n']:,}"])
    add_table(doc, ["模型" if lang == "cn" else "모형", "系数" if lang == "cn" else "계수", "标准误" if lang == "cn" else "표준오차", "p值" if lang == "cn" else "p값", "FDR", "N"], car_rows, [3300, 1300, 1300, 1000, 1000, 1460], font, east, 8.8)
    if (FIGURES / "02_最终模型系数图.png").exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(FIGURES / "02_最终模型系数图.png"), width=Inches(6.3))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = caption.add_run("图2 首选模型系数与95%置信区间" if lang == "cn" else "그림 2 선호모형 계수와 95% 신뢰구간")
        set_run_font(rr, font, east, 9, italic=True, color=MUTED)
    if lang == "cn":
        add_para(doc, "从年份固定效应升级到事件日固定效应与双向聚类后，负向系数的显著性减弱，说明同日共同冲击和误差相关是实质问题。KOSPI与KOSDAQ子样本方向均为负但均不显著，市场交互项的联合检验原始p=0.027，却没有在多重检验后形成可确认的市场差异。全量并发公告审计进一步显示，剔除高相关公告后p=0.192，剔除同日其他公告后p=0.523，负向点估计保留但精度和稳健性不足。", font, east)
        add_para(doc, "经济含义上，数字化战略提高1个标准差与三日CAR低约0.269个百分点相关。韩国市场中年报披露集中、企业集团关系、数字项目投资规模与兑现周期可能使投资者把更强数字行动同时视为机会和执行风险。由于安慰剂失败，这一解释只能作为与数据一致的机制，而不是已证明原因。", font, east)
    else:
        add_para(doc, "연도 고정효과에서 사건일 고정효과와 이중군집으로 강화하면 음(-)의 계수 유의성이 약해진다. 이는 동일일 공통충격과 오차상관이 실질적인 문제임을 뜻한다. KOSPI와 KOSDAQ 하위표본은 모두 음(-)의 방향이나 비유의적이며 시장상호작용 결합검정의 원 p=0.027도 다중검정 후 확인적 시장차이로 남지 않는다. 전수 동시공시 감사에서 고관련 공시 제외 후 p=0.192, 동일일 다른 공시 제외 후 p=0.523으로 음(-)의 점추정치는 남지만 정밀성과 강건성이 부족하다.", font, east)
        add_para(doc, "경제적 크기는 디지털 전략 1표준편차 증가가 약 0.269%p 낮은 3일 CAR과 연관됨을 의미한다. 한국시장에서 연차공시의 집중, 기업집단 관계, 디지털 사업의 투자규모와 회수기간은 투자자로 하여금 강한 디지털 행동을 기회와 실행위험으로 동시에 평가하게 할 수 있다. 그러나 위약검정이 실패하므로 이는 자료와 양립 가능한 설명일 뿐 입증된 원인이 아니다.", font, east)


def add_concurrent_robustness(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("8. 事件窗内并发公告审计" if lang == "cn" else "8. 사건창 동시공시 감사", level=1)
    s = INPUTS["concurrent_summary"]
    if lang == "cn":
        add_para(doc, f"本审计不是单纯文字叙述，而是对事件研究识别的可观测混杂检验。程序按年报披露交易日至后两个交易日，对KOSPI和KOSDAQ历史公告页进行{s['query_count']:,}个日期×市场查询并逐页抓取，再用DART corp_code匹配样本企业。目标年报rcept_no本身被排除，其他修订年报、业绩、融资、治理、审计、诉讼和公平披露均保留。", font, east)
    else:
        add_para(doc, f"본 감사는 단순한 서술이 아니라 사건연구의 관찰가능 혼재사건 검정이다. 사업보고서 공시거래일부터 이후 두 거래일까지 KOSPI와 KOSDAQ 과거 공시페이지에 대해 {s['query_count']:,}개의 날짜×시장 조회를 페이지별로 수집하고 DART corp_code로 표본기업과 연결하였다. 대상 사업보고서 rcept_no만 제외하고 정정보고서, 실적, 자금조달, 지배구조, 감사, 소송과 공정공시는 모두 유지하였다.", font, east)
    audit_rows = [
        ["目标年报收件号覆盖率" if lang == "cn" else "대상 보고서 수신번호 포괄률", f"{100*s.get('focal_annual_report_receipt_coverage', 0):.2f}%"],
        ["存在任何并发公告" if lang == "cn" else "동시공시 존재", f"{s['events_with_any_concurrent']:,} ({100*s['share_with_any_concurrent']:.2f}%)"],
        ["高相关并发公告" if lang == "cn" else "고관련 동시공시", f"{s['events_with_high_relevance']:,} ({100*s['share_with_high_relevance']:.2f}%)"],
        ["同日并发公告" if lang == "cn" else "동일일 동시공시", f"{s['events_with_same_day']:,} ({100*s['share_with_same_day']:.2f}%)"],
        ["匹配的独立公告" if lang == "cn" else "매칭된 고유공시", f"{s['matched_unique_filings']:,}"],
    ]
    add_table(doc, ["指标" if lang == "cn" else "지표", "结果" if lang == "cn" else "결과"], audit_rows, [4800, 4560], font, east, 9.5)
    concurrent_labels_kr = {
        "CONCURRENT_FULL_AlignedCAR_0_p2_w": "엄격 정렬 전체 사건표본",
        "CONCURRENT_VERIFIED_FULL_AlignedCAR_0_p2_w": "대상 사업보고서 접수번호가 DART 과거공시에서 확인된 사건만",
        "CONCURRENT_CLEAN_ANY_AlignedCAR_0_p2_w": "포괄범위 확인 후 사건창 내 다른 DART 공시를 모두 제외",
        "CONCURRENT_CLEAN_HIGH_AlignedCAR_0_p2_w": "포괄범위 확인 후 사건창 내 고관련 공시 제외",
        "CONCURRENT_CLEAN_SAMEDAY_AlignedCAR_0_p2_w": "포괄범위 확인 후 사건일 당일의 다른 공시 제외",
    }
    cc_rows = [
        [
            x["description"] if lang == "cn" else concurrent_labels_kr.get(x["model_id"], x["description"]),
            fmt_coef(x),
            f"{x['std_error']:.6f}",
            fmt_p(x["p_value"]),
            fmt_p(x.get("fdr_concurrent_family")),
            f"{x['n']:,}",
        ]
        for x in INPUTS["concurrent_models"]
    ]
    add_table(doc, ["样本" if lang == "cn" else "표본", "系数" if lang == "cn" else "계수", "标准误" if lang == "cn" else "표준오차", "p值" if lang == "cn" else "p값", "FDR", "N"], cc_rows, [3500, 1200, 1200, 1000, 1000, 1460], font, east, 8.8)
    if lang == "cn":
        add_para(doc, "正式判断以剔除任何并发公告、剔除高相关公告和剔除同日公告后系数的方向、幅度和置信区间是否保持为准，而不是只观察某一列是否跨过0.05。并发公告审计降低了把其他重大信息错误归因于数字战略文本的风险，但无法排除媒体报道、分析师活动或报告之外的信息。", font, east)
    else:
        add_para(doc, "공식 판단은 모든 동시공시 제외, 고관련 공시 제외 및 동일일 공시 제외 후 계수의 방향·크기·신뢰구간이 유지되는지를 기준으로 하며 한 열의 p값이 0.05를 넘는지만 보지 않는다. 동시공시 감사는 다른 중요정보를 디지털 전략 텍스트에 잘못 귀속하는 위험을 줄이지만 언론보도, 애널리스트 활동 또는 보고서 밖 정보까지 제거하지는 못한다.", font, east)

    doc.add_heading("9. 非线性、机制与稳健性" if lang == "cn" else "9. 비선형성·메커니즘·강건성", level=1)
    joint = INPUTS["joint_tests"]
    min_joint = min(x["p_value"] for x in joint if x.get("p_value") is not None)
    if lang == "cn":
        add_para(doc, f"经营路径的两部分、分位组、分布滞后、内部成分和正值样本三次项联合检验均未支持稳定的非线性或机制关系；相关联合检验p值最低仍高于0.20。资本市场的市场交互联合检验原始p={min_joint:.3f}，但分市场边际效应均不显著，且整个CAR模型族FDR不支持确认性结论。", font, east)
        add_para(doc, "共线性不是主要障碍：核心、扩展与治理控制的最大VIF约为2.64、3.17和3.32。更大的问题是扩展控制缺失导致样本锐减、数字指标高持久性和零值集中。未来一年数字化水平显著预测当前CAR（p=0.0035），而未来变化量不显著，提示稳定公司特征、披露风格或残余时间错配仍在驱动水平结果。", font, east)
    else:
        add_para(doc, f"경영경로의 허들, 분위집단, 분포시차, 내부성분과 양(+)값 표본 3차항 결합검정은 안정적인 비선형 또는 메커니즘 관계를 지지하지 않으며 관련 p값의 최솟값도 0.20보다 크다. 자본시장의 시장상호작용 결합검정 원 p={min_joint:.3f}이나 시장별 한계효과는 모두 비유의적이고 전체 CAR 모형군의 FDR도 확인적 결론을 지지하지 않는다.", font, east)
        add_para(doc, "다중공선성은 핵심 장애가 아니다. 핵심·확장·지배구조 통제의 최대 VIF는 약 2.64, 3.17, 3.32이다. 더 큰 문제는 확장통제의 결측으로 인한 표본감소, 디지털 지표의 높은 지속성과 영(0)값 집중이다. 미래 1년 디지털 수준이 현재 CAR을 유의하게 예측(p=0.0035)하지만 미래 변화량은 비유의적이어서 안정적 기업특성, 공시스타일 또는 잔여 시점오차가 수준결과를 유도할 가능성이 있다.", font, east)


def add_contextual_interpretation(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("10. 原因解释：为何严格模型不显著" if lang == "cn" else "10. 원인 해석: 엄격한 모형에서 유의성이 사라지는 이유", level=1)
    if lang == "cn":
        doc.add_heading("10.1 这种结果究竟说明什么", level=2)
        add_para(doc, "本研究估计的不是‘数字化是否有用’这一宽泛命题，而是同一家公司年报中实质数字行动披露相对自身历史水平的变化，在扣除公司固定特征、年度共同冲击和一组财务状态后，能否稳定预测下一年度经营结果或披露日附近的异常收益。因此，不显著的准确含义是：年报披露强度没有表现出稳定的短期增量预测力，而不是企业数字化的真实经济价值等于零。", font, east)
        add_para(doc, "这一结果更接近‘战略存量与年度流量不匹配’。数字基础设施、数据治理、流程再造和商业模式迁移通常跨越多个年度，但年度文本指标记录的是某一报告期内可观察到的披露变化。数字指标在公司内高度持续且51.56%的公司年度为零，使公司固定效应模型主要依赖较少的年度内变动；如果真实价值来自长期累积、跨部门整合或少数高质量项目，平均短期系数自然可能接近零。", font, east)
        add_para(doc, "双轮驱动也未必同步。内部路径可能先出现投资、培训、系统迁移和组织摩擦，随后才形成效率或收入；外部路径则可能同时包含增长机会与实施风险。因而‘内部短期不显著、外部弱负向’与延迟兑现和信号含义不确定相一致，但由于安慰剂和多重检验没有通过，这只能作为机制解释，不能写成已经证实的原因。", font, east)

        doc.add_heading("10.2 韩国市场与政策环境为何可能放大这种现象", level=2)
        add_para(doc, "第一，年报是低频、集中且信息捆绑的披露载体。韩国DART规定事业报告原则上在财年结束后90日内提交，大量12月结账公司由此集中在次年3月附近披露。数字战略信息可能已通过业绩预告、IR材料、新闻或项目公告被市场提前吸收；正式年报事件日又同时包含财务、审计、治理和修订信息。本研究发现64.75%的事件窗存在其他DART公告，使用精确事件日固定效应并剔除并发公告后显著性进一步减弱，说明三日CAR很难被解释为纯粹的数字战略定价。", font, east)
        add_para(doc, "第二，韩国资本市场不是简单意义上的‘更差’，而是所有权结构和估值约束不同。OECD《公司治理事实手册2025》显示，韩国上市股权中公司持有比例约27%，高于OECD成员约6%；机构投资者比例约17%，低于OECD成员约58%。结合企业集团控制、关联关系和少数股东保护问题，投资者可能更重视可验证的现金流、资本效率和股东回报，而不是孤立的战略叙述。韩国金融委员会2024年Value-up计划也把ROE、ROIC、估值、治理、沟通和股东回报作为改善‘Korea discount’的核心。这支持一种谨慎推论：没有预算、里程碑和价值兑现路径的数字化语言，其边际定价权可能有限；但本研究没有直接检验投资者结构，因此不能把所有权差异写成已确认的因果机制。", font, east)
        add_para(doc, "第三，政策普及可能制造共同趋势而不是清晰处理差异。制造业创新3.0早于样本起点，数字新政在2020年全国实施并与COVID-19重合；这些政策提高全市场数字化讨论和采用概率，却缺少未受影响的国内对照组。年份固定效应和行业×年份相对指标会吸收这种共同上升，这正是避免把宏观热潮误当企业效应所必需的。由此得到的不显著不能解释为政策失败，只能说明当前数据无法把普遍政策背景与企业特有冲击分离。", font, east)
        add_para(doc, "第四，韩国智能制造存在明显成熟度和适用范围差异。韩国研究显示，实际智能制造成熟度较高的163家制造业中小企业可能获得更好的财务绩效和运营效率，但效果取决于行业与生产人员结构；另一项韩国研究指出，2019年约75%的智能工厂仍处于基础水平，基础与中间一级合计约95%。因此，针对实际采用程度、制造企业和具体产线的研究可能显著，而覆盖KOSPI与KOSDAQ多行业上市公司的年报文字指标会把高成熟项目、初级采用、服务业平台项目和象征性披露平均在一起，导致平均效应被稀释。", font, east)

        doc.add_heading("10.3 为什么一些既有论文显著，而本研究不显著", level=2)
        add_para(doc, "首先，解释变量不同。既有显著研究常使用实际IT资本、系统成熟度、政府受益名单、问卷采用程度或工厂级技术，而本研究使用年报中可观察到的战略行动披露。前者更接近真实处理，后者同时含有实施、选择性披露和语言测量误差。Bloom、Sadun与Van Reenen（2012）发现IT生产率优势主要与管理实践等组织互补资产有关，说明技术本身或技术表述并不足以保证绩效。", font, east)
        add_para(doc, "其次，样本、结果变量和时间窗不同。韩国智能工厂论文通常研究制造业中小企业、受支持企业或工厂，结果是生产率、运营效率、采用概率或较长时期的财务变化；本研究覆盖多行业上市公司，检验下一年会计绩效和三日CAR。若收益需要两至五年实现，或只在高成熟制造项目中出现，一年平均效应和短事件窗都会偏弱。", font, east)
        add_para(doc, "再次，识别强度不同。本研究加入公司固定效应、年份或精确事件日固定效应、公司或双向聚类标准误、缩尾、逐年剔除、并发公告排除、安慰剂和FDR校正。较简单的横截面、普通OLS、单一模型或未处理共同日期冲击的研究更容易获得传统p<0.05。Hanelt等（2021）的系统综述也强调数字化转型是受组织、行业、国家和生态系统边界条件共同塑造的持续组织变革，因此跨设计复制不应预期得到相同系数。", font, east)
        add_para(doc, "最后，研究问题实际上不同。其他论文的显著结论可以说明‘在特定企业、达到一定成熟度并具备互补能力时，真实数字采用可能改善绩效’；本研究的不显著说明‘仅凭年度报告中的相对披露强度，不能在韩国上市公司总体中稳定识别短期增量回报’。两者并不矛盾，反而共同指向成熟度、组织能力、处理强度和兑现周期这些边界条件。", font, east)
        add_callout(doc, "正式讨论结论", "不能写成‘韩国市场比其他市场差’或‘韩国政策无效’。可以写成：韩国的集中披露制度、所有权与估值环境、全国性数字政策和智能制造成熟度差异，降低了年报数字战略文字作为短期边际信号的可识别度；真实采用的长期和条件性价值仍需企业级处理名单与动态模型检验。", font, east, fill=WARNING)
    else:
        doc.add_heading("10.1 이 결과가 실제로 의미하는 것", level=2)
        add_para(doc, "본 연구가 추정하는 것은 ‘디지털화가 유용한가’라는 포괄적 명제가 아니다. 동일 기업의 사업보고서에서 실질 디지털 행동공시가 과거 자기수준보다 얼마나 변했는지가 기업 고유특성, 연도 공통충격과 재무상태를 통제한 뒤 차기 경영성과 또는 공시일 주변 초과수익률을 안정적으로 예측하는지를 추정한다. 따라서 비유의성의 정확한 의미는 사업보고서 공시강도에 안정적인 단기 추가예측력이 없다는 것이며 디지털화의 실제 경제가치가 0이라는 뜻이 아니다.", font, east)
        add_para(doc, "이 결과는 ‘전략의 축적량과 연간 변화량의 불일치’에 가깝다. 디지털 인프라, 데이터 거버넌스, 프로세스 재설계와 비즈니스모델 전환은 여러 해에 걸쳐 진행되지만 연간 텍스트 지표는 특정 보고기간에 관찰된 공시변화를 기록한다. 디지털 지표의 기업내 지속성이 높고 기업-연도의 51.56%가 0이므로 기업 고정효과 모형은 제한된 연도내 변화에 의존한다. 실제 가치가 장기축적, 부서간 통합 또는 소수 고품질 프로젝트에서 발생한다면 평균 단기계수는 0에 가까울 수 있다.", font, east)
        add_para(doc, "이중 구동 경로도 동시에 작동할 필요가 없다. 내부경로에서는 투자, 교육, 시스템 이전과 조직마찰이 먼저 나타나고 효율 또는 수익은 나중에 발생할 수 있다. 외부경로는 성장기회와 실행위험을 동시에 전달할 수 있다. 따라서 ‘내부 단기 비유의, 외부 약한 음(-)’은 실현지연과 신호의 모호성에 부합하지만 위약검정과 다중검정을 통과하지 못했으므로 검증된 원인이 아니라 메커니즘 해석으로만 제시해야 한다.", font, east)

        doc.add_heading("10.2 한국 시장과 정책환경이 이 현상을 강화할 수 있는 이유", level=2)
        add_para(doc, "첫째, 사업보고서는 저빈도·집중형·정보결합형 공시수단이다. DART에 따르면 사업보고서는 원칙적으로 회계연도 종료 후 90일 이내 제출되므로 12월 결산기업의 공시가 다음해 3월 부근에 집중된다. 디지털 전략정보는 실적예고, IR자료, 뉴스 또는 사업공시를 통해 이미 반영될 수 있고 사업보고서 사건일에는 재무, 감사, 지배구조와 정정정보가 함께 나타난다. 본 연구에서 사건창의 64.75%에 다른 DART 공시가 있었고 정확한 사건일 고정효과와 동시공시 제외 후 유의성이 약해졌다. 따라서 3일 CAR을 순수 디지털 전략가격으로 해석하기 어렵다.", font, east)
        add_para(doc, "둘째, 한국 자본시장이 단순히 ‘열등하다’기보다 소유구조와 가치평가 제약이 다르다. OECD 기업지배구조 팩트북 2025에 따르면 한국 상장주식의 기업 보유비중은 약 27%로 OECD 회원국 약 6%보다 높고 기관투자자 비중은 약 17%로 OECD 회원국 약 58%보다 낮다. 기업집단 통제, 계열관계와 일반주주 보호문제를 고려하면 투자자는 독립적 전략서술보다 검증 가능한 현금흐름, 자본효율과 주주환원을 더 중시할 수 있다. 금융위원회의 2024년 밸류업 계획도 ROE, ROIC, 가치평가, 지배구조, 소통과 주주환원을 ‘코리아 디스카운트’ 개선의 핵심으로 제시한다. 이는 예산, 이정표와 가치실현 경로가 없는 디지털 언어의 한계정보가 작을 수 있다는 신중한 추론을 지지하지만 본 연구가 투자자구조를 직접 검정하지 않았으므로 확인된 인과메커니즘으로 쓸 수는 없다.", font, east)
        add_para(doc, "셋째, 정책확산은 명확한 처리차이보다 공통추세를 만들 수 있다. 제조업 혁신 3.0은 표본 시작 이전이고 2020년 디지털 뉴딜은 전국적으로 시행되면서 COVID-19와 중첩된다. 이 정책들은 시장전체의 디지털 담론과 채택확률을 높이지만 국내 비처리 대조집단을 제공하지 않는다. 연도 고정효과와 산업×연도 상대지표가 이러한 공통상승을 흡수하는 것은 거시적 유행을 기업효과로 오인하지 않기 위해 필요하다. 따라서 비유의성을 정책실패로 해석할 수 없고 현재자료가 보편적 정책배경과 기업고유 충격을 분리하지 못한다는 뜻이다.", font, east)
        add_para(doc, "넷째, 한국 스마트제조는 성숙도와 적용범위의 차이가 크다. 한국 제조업 중소기업 163개를 대상으로 한 연구는 실제 스마트제조 성숙도가 높을수록 재무성과와 운영효율이 개선될 수 있지만 효과가 산업과 생산인력구성에 따라 달라짐을 보여준다. 다른 한국 연구는 2019년 스마트공장의 약 75%가 기초수준이고 기초와 중간1 수준의 합계가 약 95%라고 보고한다. 실제 도입정도, 제조기업과 생산라인을 대상으로 한 연구는 유의할 수 있지만 KOSPI와 KOSDAQ의 다산업 상장기업 사업보고서 지표는 고성숙 프로젝트, 초기도입, 서비스 플랫폼과 상징적 공시를 평균하여 평균효과를 희석할 수 있다.", font, east)

        doc.add_heading("10.3 일부 선행연구는 유의하지만 본 연구는 비유의한 이유", level=2)
        add_para(doc, "첫째, 설명변수가 다르다. 유의한 선행연구는 실제 IT자본, 시스템 성숙도, 정부 수혜목록, 설문 채택정도 또는 공장기술을 사용하는 경우가 많다. 본 연구는 사업보고서에서 관찰되는 전략행동 공시를 사용한다. 전자는 실제 처리에 더 가깝지만 후자는 실행, 선택적 공시와 언어측정오차를 함께 포함한다. Bloom, Sadun과 Van Reenen(2012)은 IT 생산성 우위가 인사관리 등 조직적 보완자산과 주로 관련됨을 보여주어 기술 또는 기술서술만으로 성과가 보장되지 않음을 시사한다.", font, east)
        add_para(doc, "둘째, 표본, 결과변수와 시간창이 다르다. 한국 스마트공장 연구는 제조업 중소기업, 지원기업 또는 공장을 대상으로 생산성, 운영효율, 채택확률이나 장기 재무변화를 분석하는 경우가 많다. 본 연구는 다산업 상장기업을 대상으로 차기 회계성과와 3일 CAR을 검정한다. 효과가 2~5년에 걸쳐 실현되거나 고성숙 제조프로젝트에만 나타난다면 1년 평균효과와 짧은 사건창은 약할 수 있다.", font, east)
        add_para(doc, "셋째, 식별강도가 다르다. 본 연구는 기업 고정효과, 연도 또는 정확한 사건일 고정효과, 기업 또는 이중군집 표준오차, 윈저라이징, 연도별 제외, 동시공시 제거, 위약검정과 FDR 보정을 적용한다. 단순 횡단면, 일반 OLS, 단일모형 또는 공통 날짜충격을 처리하지 않은 연구는 전통적 p<0.05를 얻기 쉽다. Hanelt 등(2021)의 체계적 문헌고찰도 디지털 전환이 조직, 산업, 국가와 생태계 경계조건이 함께 형성하는 지속적 조직변화임을 강조하므로 서로 다른 설계에서 동일계수를 기대해서는 안 된다.", font, east)
        add_para(doc, "마지막으로 연구질문 자체가 다르다. 다른 논문의 유의한 결과는 ‘특정 기업이 충분한 성숙도와 보완역량을 갖추고 실제 디지털 기술을 도입하면 성과가 개선될 수 있다’는 것을 보여줄 수 있다. 본 연구의 비유의 결과는 ‘연차보고서의 상대적 공시강도만으로 한국 상장기업 전체의 단기 추가수익을 안정적으로 식별할 수 없다’는 뜻이다. 두 결과는 모순되지 않으며 성숙도, 조직역량, 처리강도와 실현기간이라는 경계조건을 함께 가리킨다.", font, east)
        add_callout(doc, "공식 논의 결론", "‘한국 시장이 다른 시장보다 나쁘다’거나 ‘한국 정책이 효과가 없다’고 쓰면 안 된다. 한국의 집중공시제도, 소유·가치평가 환경, 전국적 디지털 정책과 스마트제조 성숙도 차이가 사업보고서 디지털 전략문장의 단기 한계신호로서 식별가능성을 낮춘다고 쓸 수 있다. 실제 도입의 장기적·조건부 가치는 기업별 처리목록과 동태모형으로 추가 검정해야 한다.", font, east, fill=WARNING)


def add_policy(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("11. 韩国制度环境与外生冲击" if lang == "cn" else "11. 한국 제도환경과 외생충격", level=1)
    if lang == "cn":
        add_para(doc, "候选政策经过时间、处理组、对照组和排除限制四项审计。制造业创新3.0在2014年公布，早于2015年样本起点，因此无法提供充分样本内处理前时期。2020年数字新政全国同时实施并与COVID-19重合，简单Post2020会把疫情、宏观刺激和数字政策混在一起。两者适合制度背景，不适合作为当前模型的简单外生冲击。", font, east)
        add_para(doc, "最有希望的是企业分期采用的智能工厂支持项目。官方与研究资料显示2014-2022年累计建设约30,144家智能工厂，主管部门微观资料可包含支持企业、采用年份、类型、水平和地区。若取得名单，应限制制造业并首先补充事业者登记号，再将企业级处理年份与DART corp_code审慎匹配。", font, east)
        add_number(doc, "先做重叠和功效审计：处理企业数、采用年份、处理前三期、行业和市场分布。", font, east)
        add_number(doc, "采用Callaway-Sant'Anna组别-时间ATT或Sun-Abraham动态事件研究，使用从未处理或尚未处理企业。", font, east)
        add_number(doc, "报告处理前趋势、预期效应、建设等级和项目类型异质性，并处理政策批次或地区共同冲击。", font, east)
        add_number(doc, "在企业级处理名单取得与核验之前，只把该设计列为预注册式后续扩展，不声称已完成因果估计。", font, east)
    else:
        add_para(doc, "후보정책은 시점, 처리집단, 대조집단과 배제제약의 네 기준으로 감사하였다. 제조업 혁신 3.0은 2014년에 발표되어 2015년 표본시작보다 앞서므로 충분한 표본내 사전기간을 제공하지 않는다. 2020년 디지털 뉴딜은 전국 동시시행이고 COVID-19와 겹쳐 단순 Post2020이 감염병, 거시부양과 디지털 정책을 혼합한다. 두 정책은 제도적 배경에는 적합하지만 현재모형의 단순 외생충격에는 부적합하다.", font, east)
        add_para(doc, "가장 유망한 것은 기업별로 순차 도입된 스마트공장 지원사업이다. 공식 및 연구자료에 따르면 2014-2022년 누적 약 30,144개의 스마트공장이 구축되었고 담당부처 미시자료는 지원기업, 도입연도, 유형, 수준과 지역을 포함할 수 있다. 목록을 확보하면 제조업으로 제한하고 사업자등록번호를 보완한 뒤 기업별 처리연도를 DART corp_code와 신중히 매칭해야 한다.", font, east)
        add_number(doc, "처리기업 수, 도입연도, 처리 전 3개 기간, 산업 및 시장분포에 대한 중첩·검정력 감사를 먼저 수행한다.", font, east)
        add_number(doc, "미처리 또는 아직 처리되지 않은 기업을 대조군으로 Callaway-Sant'Anna 그룹-시간 ATT 또는 Sun-Abraham 동태사건연구를 사용한다.", font, east)
        add_number(doc, "사전추세, 예상효과, 구축수준과 사업유형 이질성을 보고하고 정책차수 또는 지역 공통충격을 처리한다.", font, east)
        add_number(doc, "기업별 처리목록을 확보·검증하기 전에는 사전등록형 후속확장으로만 제시하고 인과추정 완료를 주장하지 않는다.", font, east)
    policy_rows = [
        ["制造业创新3.0" if lang == "cn" else "제조업 혁신 3.0", "2014", "背景，不用于简单DiD" if lang == "cn" else "배경, 단순 DiD 제외"],
        ["数字新政" if lang == "cn" else "디지털 뉴딜", "2020", "全国冲击且与COVID重合" if lang == "cn" else "전국충격 및 COVID 중첩"],
        ["智能工厂分期采用" if lang == "cn" else "스마트공장 순차도입", "2014-", "取得微观受益名单后可做现代DiD" if lang == "cn" else "미시 수혜목록 확보 후 현대적 DiD"],
    ]
    add_table(doc, ["政策" if lang == "cn" else "정책", "时间" if lang == "cn" else "시점", "识别判断" if lang == "cn" else "식별판단"], policy_rows, [2800, 1400, 5160], font, east, 9.2)


def add_conclusion(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("12. 结论、贡献与限制" if lang == "cn" else "12. 결론·기여·한계", level=1)
    if lang == "cn":
        add_para(doc, "第一，内部真实效应的正式结论是没有发现稳健的短期经营绩效改善。第二，外部资本市场的正式结论是更强的实质数字战略披露与较低三日CAR存在弱负向关联，但该关系没有通过多重检验，且受到安慰剂失败约束。第三，非线性、BMI重构和纯市场信号没有形成可确认的独立机制。", font, east)
        add_para(doc, "跨国比较不能概括为韩国市场‘更差’。更有依据的表述是：韩国年报披露时点集中、公司与机构投资者结构不同、企业集团与股东价值问题更受关注、全国性数字政策形成共同趋势，且智能制造成熟度高度异质。这些制度特征降低了年报数字战略语言作为短期边际信号的可识别度，却不否定真实数字采用在具备组织互补资产和较长兑现周期时可能创造价值。", font, east)
        add_para(doc, "论文的理论意义不是宣称数字化必然提升绩效，而是区分战略行动、商业模式改变、可见信号和口号性语言，并展示披露、实施与价值实现之间可能存在时间错位。对韩国市场而言，集中披露日、数字政策普及、KOSPI/KOSDAQ结构差异和执行不确定性使简单词频与简单事件研究尤其容易产生误判。", font, east)
        add_para(doc, "方法贡献包括：大规模韩文年报自动采集；句子级行动筛选；行业×年份相对信号；修订文本与事件日严格对齐；公司与事件日双向聚类；并发公告逐日逐页审计；零膨胀、持久性、非线性、多重检验和时间边界的系统诊断。负面或弱结果在这一设计下具有信息价值，因为它显示文献中的显著关系可能依赖粗糙测量、选择性样本、共同趋势或过于乐观的标准误。", font, east)
        add_para(doc, "主要限制是文本指标仍不能直接测量投入金额、项目质量与组织能力；英文披露不足迫使研究转向韩文年报；独立第二编码者验证仍需真实完成；财务与治理扩展变量存在缺失；事件研究无法排除所有非DART信息；当前政策冲击缺少企业级处理名单。因此当前版本达到的是严谨的关联性实证设计，而不是最终的因果论文。", font, east)
    else:
        add_para(doc, "첫째, 내부 실질효과의 공식결론은 강건한 단기 경영성과 개선을 발견하지 못했다는 것이다. 둘째, 외부 자본시장에서 강한 실질 디지털 전략 공시는 낮은 3일 CAR과 약한 음(-)의 연관성을 보이나 다중검정을 통과하지 못하고 위약검정 실패의 제약을 받는다. 셋째, 비선형성, BMI 재구성과 순수시장신호는 확인 가능한 독립 메커니즘을 형성하지 못했다.", font, east)
        add_para(doc, "국가간 비교를 한국 시장이 ‘더 나쁘다’고 요약할 수는 없다. 보다 근거있는 표현은 한국 사업보고서 공시시점의 집중, 기업 및 기관투자자 구성의 차이, 기업집단과 주주가치 문제에 대한 높은 관심, 전국적 디지털정책의 공통추세와 스마트제조 성숙도의 이질성이 사업보고서 디지털 전략언어의 단기 한계신호로서 식별가능성을 낮춘다는 것이다. 이는 조직적 보완자산과 긴 실현기간을 갖춘 실제 디지털 도입의 가치 가능성을 부정하지 않는다.", font, east)
        add_para(doc, "이론적 의의는 디지털화가 반드시 성과를 높인다고 주장하는 데 있지 않다. 전략행동, 비즈니스모델 변화, 가시적 신호와 상투적 언어를 구분하고 공시·실행·가치실현 사이의 시간불일치를 제시한다. 한국시장에서는 집중된 공시일, 디지털정책 확산, KOSPI/KOSDAQ 구조차이와 실행불확실성 때문에 단순 키워드 빈도와 단순 사건연구가 특히 오판을 만들 수 있다.", font, east)
        add_para(doc, "방법론적 기여는 대규모 한국어 사업보고서 자동수집, 문장수준 행동필터, 산업×연도 상대신호, 정정텍스트와 사건일의 엄격한 정렬, 기업·사건일 이중군집, 동시공시의 일별·페이지별 감사, 영(0)값·지속성·비선형성·다중검정·시점경계의 체계적 진단이다. 이러한 설계의 비유의 또는 약한 결과는 기존 유의관계가 거친 측정, 선택표본, 공통추세 또는 낙관적 표준오차에 의존할 수 있음을 보여준다는 점에서 정보가치가 있다.", font, east)
        add_para(doc, "주요 한계는 텍스트지표가 투자금액, 사업품질과 조직역량을 직접 측정하지 못하고 영문공시 부족으로 한국어 보고서에 의존한다는 점이다. 실제 독립 제2코더 검증이 아직 필요하고 재무·지배구조 확장변수에 결측이 있으며 사건연구가 모든 비DART 정보를 제거하지 못한다. 현재 정책충격에는 기업별 처리목록도 없다. 따라서 현 버전은 엄격한 연관성 실증설계이지 완성된 인과논문은 아니다.", font, east)
    add_callout(doc, "投稿口径" if lang == "cn" else "투고 표현", "使用“未发现稳健证据”“与……相关”“结果与实施成本/兑现风险解释一致”，避免使用“证明”“导致”“惩罚”或“中介效应成立”。" if lang == "cn" else "‘강건한 증거를 발견하지 못함’, ‘…와 연관됨’, ‘실행비용/실현위험 해석과 양립함’이라고 표현하고 ‘입증’, ‘원인’, ‘페널티’ 또는 ‘매개효과 성립’을 피한다.", font, east, fill=WARNING)


def add_repro_refs(doc: Document, lang: str, font: str, east: str) -> None:
    doc.add_heading("13. 复现与文件说明" if lang == "cn" else "13. 재현 및 파일안내", level=1)
    if lang == "cn":
        add_para(doc, "正式包的01_核心数据包含冻结分析面板、严格对齐事件研究数据、事件级并发公告标记、并发公告明细和变量字典。02_模型与结果包含总览工作簿、完整经营与CAR模型、联合检验、逐年剔除、文本时间边界和并发公告剔除模型。06_复现代码保留最终模型、事件审计和模型审计脚本。", font, east)
        add_para(doc, "复现顺序为：读取冻结面板并构造派生变量；运行经营与CAR模型；进行FDR、VIF、零膨胀与时间边界诊断；抓取DART事件窗公告并重跑清洁样本；最后生成正式资产、工作簿和双语报告。API密钥只从环境变量读取，不写入代码或成果包。", font, east)
    else:
        add_para(doc, "공식 패키지의 핵심데이터 폴더에는 동결분석패널, 엄격 정렬 사건자료, 사건별 동시공시 표식, 동시공시 상세와 변수사전이 있다. 모형·결과 폴더에는 요약 통합문서, 전체 경영 및 CAR 모형, 결합검정, 연도별 제외, 텍스트 시점경계와 동시공시 제외모형이 포함된다. 재현코드 폴더에는 최종모형, 사건감사와 모형감사 스크립트가 있다.", font, east)
        add_para(doc, "재현순서는 동결패널에서 파생변수 구축, 경영 및 CAR 모형 추정, FDR·VIF·영(0)값·시점경계 진단, DART 사건창 공시수집과 정제표본 재추정, 최종 자산·통합문서·이중언어 보고서 생성이다. API 키는 환경변수에서만 읽고 코드나 결과패키지에 기록하지 않는다.", font, east)

    doc.add_heading("14. 参考来源" if lang == "cn" else "14. 참고자료", level=1)
    for source in INPUTS["sources"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(5)
        text = f"{source['organization']} ({source['year']}). {source['title']}. {source['url']}"
        run = p.add_run(text)
        set_run_font(run, font, east, 9.5, color=TEXT)


def build(lang: str, output: Path) -> None:
    doc, font, east = setup_document(lang)
    add_cover(doc, lang, font, east)
    add_contents(doc, lang, font, east)
    add_executive_summary(doc, lang, font, east)
    add_theory(doc, lang, font, east)
    add_data_measurement(doc, lang, font, east)
    add_models(doc, lang, font, east)
    add_results(doc, lang, font, east)
    add_concurrent_robustness(doc, lang, font, east)
    add_contextual_interpretation(doc, lang, font, east)
    add_policy(doc, lang, font, east)
    add_conclusion(doc, lang, font, east)
    add_repro_refs(doc, lang, font, east)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    cn = FINAL / "03_中文报告" / "企业数字化战略双轮驱动效应_正式实证报告_中文.docx"
    kr = FINAL / "04_韩文报告" / "기업_디지털_전략의_이중_구동_효과_공식_실증보고서_한국어.docx"
    build("cn", cn)
    build("kr", kr)
    print(json.dumps({"chinese": str(cn), "korean": str(kr)}, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
