from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
FIGURES = ROOT / "figures"
OUTPUT = ROOT / "report" / "Ames房价预测_实证分析报告.docx"

INK = "1E2933"
BLUE = "2F6B8A"
GREEN = "2F7D67"
ORANGE = "D07A32"
LIGHT_BLUE = "EAF2F6"
LIGHT_GREEN = "E9F3EF"
LIGHT_GRAY = "F2F4F5"
MID_GRAY = "66717D"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text_color(cell, color: str) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    paragraph.add_run(" 页")


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 27, INK, 0, 10),
        ("Subtitle", 13, MID_GRAY, 0, 16),
        ("Heading 1", 17, BLUE, 18, 8),
        ("Heading 2", 13, GREEN, 12, 5),
        ("Heading 3", 11, ORANGE, 8, 3),
    ):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.font.italic = False
    caption.paragraph_format.space_after = Pt(10)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("AMES HOUSE PRICE PREDICTION")
    set_run_font(run, size=8.5)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph.add_run("    技术报告 · 2026").font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_accent_rule(document: Document, color: str = GREEN) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.6)
    cell = table.cell(0, 0)
    cell.height = Cm(0.13)
    set_cell_shading(cell, color)
    cell.text = ""


def add_cover(document: Document, summary: dict, rank: pd.Series) -> None:
    document.add_paragraph("数据分析与机器学习作品集", style="Subtitle")
    title = document.add_paragraph(style="Title")
    title.add_run("Ames 房价预测")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("从基线筛选到稳健四模型融合")
    add_accent_rule(document)
    document.add_paragraph(
        "围绕数据语义、交叉验证、模型互补性和误差边界完成的一次端到端回归实践。"
    )

    metrics = document.add_table(rows=2, cols=4)
    metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
    metrics.autofit = False
    labels = ["Kaggle Public", "排名快照", "排名百分比", "稳健本地 RMSE"]
    values = [
        "0.12163",
        f"{int(rank['Rank'])} / {int(rank['TotalTeams']):,}",
        f"Top {float(rank['TopPercent']):.2f}%",
        f"{float(summary['robust_crossfit_calibrated_rmse']):.5f}",
    ]
    for index, label in enumerate(labels):
        metrics.cell(0, index).text = label
        metrics.cell(1, index).text = values[index]
        set_cell_shading(metrics.cell(0, index), LIGHT_BLUE)
        set_cell_shading(metrics.cell(1, index), LIGHT_GREEN)
        metrics.cell(0, index).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        metrics.cell(1, index).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in metrics.cell(0, index).paragraphs[0].runs:
            set_run_font(run, size=8.5)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BLUE)
        for run in metrics.cell(1, index).paragraphs[0].runs:
            set_run_font(run, size=14)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(INK)
        metrics.cell(0, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        metrics.cell(1, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("")
    paragraph = document.add_paragraph()
    paragraph.add_run("最终结论  ").bold = True
    paragraph.add_run(
        "V3 比 V2 的公开分数改善 0.00042，排名快照由 493 升至 430。"
        "结果可复现，但本地验证仍明显偏乐观，价格两端仍是主要误差来源。"
    )
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(20)
    run = note.add_run("结果快照：2026-08-24  |  Kaggle team: kimzaiyi1")
    set_run_font(run, size=9)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_callout(document: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    label = paragraph.add_run(title + "  ")
    label.bold = True
    label.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph.add_run(body)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = widths is None
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        header.cells[index].text = text
        set_cell_shading(header.cells[index], BLUE)
        set_cell_text_color(header.cells[index], WHITE)
        header.cells[index].paragraphs[0].runs[0].bold = True
        header.cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = str(value)
            cells[column_index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cells[column_index], LIGHT_GRAY)
        if widths:
            for cell, width in zip(cells, widths):
                cell.width = Cm(width)
    document.add_paragraph("")


def add_figure(document: Document, filename: str, caption: str, width: float = 6.45) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build() -> None:
    summary = json.loads((RESULTS / "v3_summary.json").read_text(encoding="utf-8"))
    rank = pd.read_csv(RESULTS / "kaggle_rank_snapshot.csv").iloc[0]
    history = pd.read_csv(RESULTS / "kaggle_submission_history.csv")
    baseline = pd.read_csv(RESULTS / "baseline_models.csv").sort_values("oof_rmse")
    residual = pd.read_csv(RESULTS / "v3_residual_by_price_decile.csv")
    reproducibility = json.loads(
        (RESULTS / "reproducibility_check.json").read_text(encoding="utf-8")
    )

    document = Document()
    document.core_properties.title = "Ames 房价预测：实证分析报告"
    document.core_properties.subject = "Kaggle House Prices 机器学习作品集"
    document.core_properties.author = "个人数据分析作品集"
    document.core_properties.keywords = "Kaggle, Ames Housing, Regression, Ensemble"
    style_document(document)
    add_cover(document, summary, rank)

    add_heading(document, "执行摘要")
    document.add_paragraph(
        "本项目使用 Ames Housing 数据预测住宅成交价格。最终方案由 Elastic Net、"
        "Gradient Boosting、XGBoost 和 RBF-SVR 组成，通过重复 OOF 预测、目标分层的"
        "元交叉验证和轻度线性校准生成提交结果。"
    )
    add_callout(
        document,
        "答案先行",
        "V3 Public Score 为 0.12163，排名快照为 430 / 3,454（Top 12.45%）。"
        "相较 V2 有真实但有限的提升。本地 0.10483 没有完全兑现到榜单，"
        "因此报告把验证偏差和尾部误差作为核心结论。",
        LIGHT_GREEN,
    )
    add_bullets(
        document,
        [
            "有效做法：结构性缺失值处理、对数目标、OOF 选型、模型互补性检查。",
            "V3 改进：加入 RBF-SVR、使用目标分层元交叉验证、限制校准幅度。",
            "主要局限：样本量小、重复实验造成选择偏差、价格两端仍回归均值。",
            "可信边界：未使用外部成交价或公开泄漏标签，不承诺隐藏测试集达到 0.06。",
        ],
    )
    add_figure(document, "09_kaggle_result.png", "图 1  Kaggle 公开分数与排名快照")

    document.add_page_break()
    add_heading(document, "1. 问题定义与数据", 1)
    document.add_paragraph(
        "训练集包含 1,460 套房屋、81 列；测试集包含 1,459 套房屋、80 列。"
        "除 Id 外共有 79 个预测变量，覆盖面积、年份、质量等级、社区和房屋类型。"
        "目标变量 SalePrice 只出现在训练集。"
    )
    add_table(
        document,
        ["数据项", "训练集", "测试集", "处理判断"],
        [
            ["行数", "1,460", "1,459", "训练时移除 2 个明确面积异常样本"],
            ["预测变量", "79", "79", "Id 不作为特征"],
            ["目标", "SalePrice", "隐藏", "预测 log(SalePrice)，提交前还原"],
            ["主要缺失", "PoolQC、Alley、Fence 等", "同类", "区分无设施与未知"],
        ],
    )
    document.add_paragraph(
        "比赛评价接近对数价格上的均方根误差。直接预测 log(SalePrice) 可以减弱少数高价房"
        "对平方误差的支配，并让误差更接近相对偏差。"
    )
    add_figure(document, "01_target_distribution.png", "图 2  原始房价右偏，对数变换后更接近对称分布")

    add_heading(document, "2. 数据预处理判断", 1)
    add_heading(document, "2.1 缺失值不是一个问题", 2)
    document.add_paragraph(
        "PoolQC、GarageType、BsmtQual 等字段中的空值经常表示房屋没有对应设施。"
        "本项目把这类值编码为 None；只有设施存在但描述字段仍为空时才标记 Missing。"
        "剩余数值缺失在每个训练折内部按中位数填充，并增加缺失指示变量。"
    )
    document.add_paragraph(
        "LotFrontage 与社区布局相关，因此 SVR 分支按训练折的 Neighborhood 中位数补充，"
        "无法匹配时再使用全局中位数。验证折没有参与统计量计算。"
    )
    add_figure(document, "02_missingness.png", "图 3  缺失最多的字段及其业务语义")
    add_heading(document, "2.2 异常值处理保持克制", 2)
    document.add_paragraph(
        "Id 524 和 1299 的居住面积超过 4,000 平方英尺，但成交价明显偏低，"
        "符合预先设定的面积异常规则，仅在模型训练阶段移除。其余高价和大面积样本保留，"
        "避免为了降低本地误差而删除真正困难的房屋。"
    )
    add_heading(document, "2.3 特征工程", 2)
    add_bullets(
        document,
        [
            "规模：总使用面积、总浴室数、总门廊面积、总室外面积。",
            "时间：售出时房龄、翻新后年限、车库年龄。",
            "设施：车库、地下室、壁炉、泳池、二层和砖石存在标记。",
            "关系：质量×居住面积、质量×总面积、车库容量×车库面积。",
            "顺序：质量、地下室采光、车库完成度等类别映射为等级分数。",
        ],
    )

    add_heading(document, "3. 模型筛选与取舍", 1)
    document.add_paragraph(
        "初始候选覆盖十类模型。所有比较使用 OOF RMSE，训练拟合误差不参与选型。"
        "最佳单模是 Lasso/Elastic Net，Gradient Boosting 和 XGBoost 单模略弱但残差互补。"
        "随机森林、Extra Trees 和早期 SVR 版本落后，因此没有为增加模型数量而保留。"
    )
    top_rows = [
        [str(row.model), f"{row.oof_rmse:.5f}", f"{row.fold_std_rmse:.5f}"]
        for row in baseline.head(7).itertuples()
    ]
    add_table(document, ["基线模型", "OOF RMSE", "折间标准差"], top_rows)
    add_figure(document, "03_baseline_model_screening.png", "图 4  十类基线模型的五折 OOF 比较")
    add_figure(document, "04_model_selection_path.png", "图 5  从最佳单模到 V3 稳健融合的验证改进")
    add_callout(
        document,
        "独立判断",
        "V3 没有继续增加树模型。V2 的树模型残差相关系数超过 0.92，"
        "再堆同类模型通常只会让本地分数更好看。重新调参后的 RBF-SVR 虽不是最强单模，"
        "但与 XGBoost 的残差相关性约为 0.853，提供了更有价值的差异。",
    )

    document.add_page_break()
    add_heading(document, "4. 最终四模型，用大白话解释", 1)
    model_rows = [
        ["Elastic Net", "给各项条件估算稳定的加减价", "稀疏高维、整体线性关系", "难表达复杂门槛"],
        ["Gradient Boosting", "后一棵树专门修正前一轮错误", "非线性、异常值稳健", "与其他树模型易同质"],
        ["XGBoost", "带复杂度约束的高效提升树", "交互、稀疏特征", "参数多、可能过拟合"],
        ["RBF-SVR", "让特征相似的房屋得到相近估价", "平滑非线性边界", "对缩放和参数敏感"],
    ]
    add_table(document, ["模型", "直观理解", "主要贡献", "主要局限"], model_rows)
    document.add_paragraph(
        "融合不是让四个模型投票选冠军，而是把它们的 OOF 预测作为新的证据。"
        "每套训练房屋都由没有见过它的模型预测，再根据这些预测估计组合权重。"
    )

    add_heading(document, "5. 权重与校准", 1)
    weights = summary["selected_weights"]
    add_table(
        document,
        ["模型", "最终权重", "权重含义"],
        [
            ["Elastic Net", f"{weights['ElasticNet_Domain']:.1%}", "稳定线性基准"],
            ["Gradient Boosting", f"{weights['GradientBoosting']:.1%}", "稳健非线性纠错"],
            ["XGBoost", f"{weights['XGBoost']:.1%}", "复杂交互与分裂"],
            ["RBF-SVR", f"{weights['SVR_RBF']:.1%}", "平滑核方法补充"],
        ],
    )
    document.add_paragraph(
        "权重要求非负、总和为 1，并加入 0.01 的 L2 惩罚，防止某个随机折把权重推向单一模型。"
        "8 组目标分层元交叉验证用于评估权重稳定性，最终使用 40 组折内权重的平均值。"
    )
    document.add_paragraph(
        f"融合后校准斜率为 {float(summary['calibration_slope']):.4f}，仅把预测范围拉开约 1.6%。"
        "斜率被限制在 0.98 至 1.04，避免根据同一批残差做激进修正。"
    )
    add_figure(document, "05_blend_weights.png", "图 6  V3 四模型融合权重")

    add_heading(document, "6. 实证结果", 1)
    result_rows = []
    for row in history.itertuples():
        local = "0.10587" if row.Version == "V2" else f"{float(summary['robust_crossfit_calibrated_rmse']):.5f}"
        result_rows.append(
            [
                row.Version,
                local,
                f"{row.PublicScore:.5f}",
                f"{int(row.RankSnapshot)} / {int(row.TotalTeams):,}",
            ]
        )
    add_table(document, ["版本", "稳健本地 RMSE", "Public Score", "排名快照"], result_rows)
    document.add_paragraph(
        "V3 预测中位数约 156,841 美元，最小值约 46,061 美元，最大值约 847,620 美元。"
        "提交文件共有 1,459 行，列名严格为 Id 和 SalePrice，Id 顺序与官方样例一致，"
        "没有缺失、重复或非正价格。"
    )
    add_figure(document, "08_prediction_distribution.png", "图 7  训练价格与测试预测价格分布")

    add_heading(document, "7. 误差诊断", 1)
    add_heading(document, "7.1 价格两端仍然最难", 2)
    first = residual.iloc[0]
    last = residual.iloc[-1]
    document.add_paragraph(
        f"最低价格十分位的 OOF RMSE 为 {float(first['rmse']):.5f}，"
        f"平均残差为 {float(first['bias']):.5f}；最高价格十分位的 OOF RMSE 为 "
        f"{float(last['rmse']):.5f}，平均残差为 {float(last['bias']):.5f}。"
        "负残差表示低价房仍被高估，正残差表示高价房仍被低估。"
    )
    add_figure(document, "06_residual_by_price_decile.png", "图 8  按实际价格十分位分组的 OOF 误差与偏差")
    add_figure(document, "07_oof_actual_vs_predicted.png", "图 9  OOF 预测与实际对数价格")
    add_heading(document, "7.2 本地验证仍偏乐观", 2)
    public_score = float(rank["Score"])
    local_score = float(summary["robust_crossfit_calibrated_rmse"])
    document.add_paragraph(
        f"V3 本地稳健 RMSE 为 {local_score:.5f}，Public Score 为 {public_score:.5f}，"
        f"差距为 {public_score - local_score:.5f}。对抗验证平均 AUC 约 0.516，"
        "接近随机判断，因此没有证据支持强烈的整体训练/测试分布漂移。"
    )
    add_bullets(
        document,
        [
            "训练样本不足 1,500 个，价格两端样本更少。",
            "多轮特征和参数实验反复使用同一训练集，产生模型选择偏差。",
            "最终四个成员的残差仍然较高相关，真正的模型多样性有限。",
            "权重和校准来自同一批 OOF 样本，不能替代完全独立的测试集。",
        ],
    )
    add_figure(document, "10_validation_gap.png", "图 10  稳健本地验证与 Kaggle Public Score 的差距")

    add_heading(document, "8. 为什么没有把 0.06 当作完成标准", 1)
    add_callout(
        document,
        "边界说明",
        "隐藏测试标签使任何具体 Kaggle 分数都无法在提交前保证。公开榜单中还存在外部标签、"
        "公开泄漏或直接重建测试答案的极低分提交，它们不能代表正常机器学习泛化能力。",
        "F8EDE8",
    )
    document.add_paragraph(
        "为了让作品经得起复现和面试追问，本项目没有使用外部成交价，也没有逐行修改测试答案。"
        "0.12163 不是接近完美，但它是一条可以解释、可以复现、没有依赖测试标签的结果。"
    )

    document.add_page_break()
    add_heading(document, "9. 质量检查与复现", 1)
    add_table(
        document,
        ["检查项", "结果", "证据"],
        [
            ["提交结构", "通过", "1,459 行；Id、SalePrice 两列"],
            ["数值有效性", "通过", "无 NaN/Inf，价格全部大于 0"],
            ["Id 一致性", "通过", "唯一、递增、与样例顺序一致"],
            ["完整重训", reproducibility["status"], f"最大 log 差 {reproducibility['max_absolute_log_difference']:.2e}"],
            ["图表", "通过", "10 张图均完成像素与尺寸检查"],
            ["Notebook", "通过", "5 个代码单元全部执行，无错误输出"],
        ],
    )
    document.add_paragraph(
        "完整重训与已提交 V3 的最大价格差约 0.15 美元，最大对数差为 "
        f"{reproducibility['max_absolute_log_difference']:.2e}。"
        "该差异来自 XGBoost 多线程浮点计算，不影响 Kaggle 五位小数分数。"
    )

    add_heading(document, "10. 下一步改进", 1)
    add_bullets(
        document,
        [
            "预留一块从未参与特征和参数选择的验证集，或使用嵌套交叉验证。",
            "针对低价和高价房设计分层模型或分位数辅助目标，并放入外层验证检验。",
            "使用重复目标分层折重训所有基础模型，比较切分策略的公共榜单稳定性。",
            "研究社区×房龄、质量×面积等分层交互，同时控制高基数类别方差。",
            "通过 bootstrap 给融合权重报告稳定区间，而不只报告一个点估计。",
        ],
    )
    add_heading(document, "结论", 1)
    document.add_paragraph(
        "项目完成了从数据理解、可复现预处理、模型筛选、融合、误差诊断到 Kaggle 提交的闭环。"
        "V3 比 V2 有小幅真实提升，排名快照进入前 12.45%。更重要的结论是："
        "后续最需要加强的并不是继续增加模型数量，而是建立更独立的验证层、"
        "降低选择偏差，并专门处理价格两端的系统性误差。"
    )

    document.add_page_break()
    add_heading(document, "附录：文件与证据索引", 1)
    add_table(
        document,
        ["内容", "位置"],
        [
            ["最终训练入口", "src/train_and_predict.py"],
            ["数据处理与公共函数", "src/house_prices_core.py"],
            ["结果复核 Notebook", "notebooks/final_results_walkthrough.ipynb"],
            ["最终提交", "results/submission_v3.csv"],
            ["Kaggle 提交历史", "results/kaggle_submission_history.csv"],
            ["排名快照", "results/kaggle_rank_snapshot.csv"],
            ["完整重训检查", "results/reproducibility_check.json"],
        ],
    )
    paragraph = document.add_paragraph()
    paragraph.add_run("口径说明：").bold = True
    paragraph.add_run(
        "分数越低越好；排名和参赛队数量为 2026-08-24 快照，后续可能变化。"
        "所有本地误差均在移除两条预定义面积异常样本后计算。"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print("Saved DOCX report")


if __name__ == "__main__":
    build()
